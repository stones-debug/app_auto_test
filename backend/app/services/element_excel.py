"""元素库 Excel 模板、导出与导入解析。"""

import json
import re
from collections.abc import Iterable
from dataclasses import dataclass
from hashlib import sha256
from io import BytesIO
from typing import Any
from zipfile import BadZipFile, ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ELEMENT_HEADERS = (
    "元素ID(element_id)",
    "项目ID(project_id)",
    "项目名称(project_name)",
    "名称(name)",
    "页面(page_name)",
    "平台(platform)",
    "适用范围(scope)",
    "定位方式(locator_type)",
    "定位值(locator_value)",
    "智能定位配置(locator_config)",
    "描述(description)",
)
ELEMENT_HEADER_FIELDS = {
    "元素ID(element_id)": "element_id",
    "项目ID(project_id)": "project_id",
    "项目名称(project_name)": "project_name",
    "名称(name)": "name",
    "页面(page_name)": "page_name",
    "平台(platform)": "platform",
    "适用范围(scope)": "scope",
    "定位方式(locator_type)": "locator_type",
    "定位值(locator_value)": "locator_value",
    "智能定位配置(locator_config)": "locator_config",
    "描述(description)": "description",
}
ELEMENT_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
CONFIG_SHEET_TITLE = "智能定位配置"
CONFIG_HEADERS = (
    "引用标识(config_ref)",
    "分片序号(chunk_index)",
    "配置内容(chunk)",
)
CONFIG_REF_PREFIX = "@config-ref:"
_CONFIG_REF_RE = re.compile(r"^@config-ref:[0-9a-f]{64}$")
EXCEL_CELL_MAX_CHARS = 32767
# 使用更保守的分片大小，兼容 Excel 对非 BMP 字符按 UTF-16 计数的实现差异。
CONFIG_CHUNK_SIZE = 16000
MAX_IMPORT_BYTES = 5 * 1024 * 1024
MAX_IMPORT_ROWS = 2000
MAX_CONFIG_SHEET_ROWS = MAX_IMPORT_ROWS * 10
MAX_EXPORT_ROWS = 20000
MAX_UNCOMPRESSED_BYTES = 50 * 1024 * 1024
MAX_ZIP_ENTRIES = 1000


@dataclass(slots=True)
class ElementImportRow:
    row_number: int
    element_id: int | None
    project_id: int | None
    data: dict[str, Any]


def _error(row: int, field: str, message: str) -> dict[str, Any]:
    return {"row": row, "field": field, "message": message}


def _text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _parse_integer(value: Any, row: int, field: str, errors: list[dict[str, Any]]) -> int | None:
    if value is None or _text(value) == "":
        return None
    try:
        number = float(value) if isinstance(value, (int, float)) else float(_text(value))
        if not number.is_integer() or number < 1:
            raise ValueError
        return int(number)
    except (TypeError, ValueError):
        errors.append(_error(row, field, "必须是大于 0 的整数"))
        return None


def _validate_archive(content: bytes) -> None:
    if len(content) > MAX_IMPORT_BYTES:
        raise ValueError("文件大小不能超过 5 MB")
    try:
        with ZipFile(BytesIO(content)) as archive:
            infos = archive.infolist()
            if len(infos) > MAX_ZIP_ENTRIES:
                raise ValueError("Excel 压缩包文件项过多")
            uncompressed = sum(info.file_size for info in infos)
            if uncompressed > MAX_UNCOMPRESSED_BYTES:
                raise ValueError("Excel 解压后大小超过限制")
    except BadZipFile as exc:
        raise ValueError("文件不是有效的 XLSX 工作簿") from exc


def _apply_data_validation(sheet) -> None:
    platform = DataValidation(type="list", formula1='"android,ios,both"', allow_blank=True)
    locator = DataValidation(
        type="list",
        formula1='"id,resource_id,xpath,accessibility_id,class_name,uiautomator,predicate,coordinate,custom,smart"',
        allow_blank=False,
    )
    sheet.add_data_validation(platform)
    sheet.add_data_validation(locator)
    platform.add(f"F2:F{MAX_IMPORT_ROWS + 1}")
    locator.add(f"H2:H{MAX_IMPORT_ROWS + 1}")


def _style_data_rows(sheet) -> None:
    # 只设置已写入的数据行；不要为了预留导入范围创建 2,000 行空单元格。
    for row in sheet.iter_rows(min_row=2, max_row=sheet.max_row):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if cell.column in (1, 2):
                cell.number_format = "0"
            else:
                cell.number_format = "@"


def _config_reference(config_text: str) -> str:
    digest = sha256(config_text.encode("utf-8")).hexdigest()
    return f"{CONFIG_REF_PREFIX}{digest}"


def _write_config_sheet(workbook: Workbook, configs: dict[str, str]) -> None:
    """Write oversized smart-locator JSON in chunks below Excel's cell limit."""
    sheet = workbook.create_sheet(CONFIG_SHEET_TITLE)
    sheet.append(list(CONFIG_HEADERS))
    for reference, config_text in configs.items():
        for index, start in enumerate(range(0, len(config_text), CONFIG_CHUNK_SIZE), start=1):
            sheet.append([reference, index, config_text[start : start + CONFIG_CHUNK_SIZE]])

    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:C{max(sheet.max_row, 2)}"
    widths = [82, 16, 100]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[chr(64 + index)].width = width
    header_fill = PatternFill("solid", fgColor="4F46E5")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.row_dimensions[1].height = 32
    for row in sheet.iter_rows(min_row=2, max_row=sheet.max_row, max_col=3):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.number_format = "0" if cell.column == 2 else "@"


def _read_config_sheet(workbook) -> tuple[dict[str, str], list[dict[str, Any]]]:
    """Read and validate optional oversized smart-locator config chunks."""
    if CONFIG_SHEET_TITLE not in workbook.sheetnames:
        return {}, []
    sheet = workbook[CONFIG_SHEET_TITLE]
    errors: list[dict[str, Any]] = []
    if sheet.max_row > MAX_CONFIG_SHEET_ROWS + 1:
        errors.append(
            _error(
                MAX_CONFIG_SHEET_ROWS + 2,
                "配置分片",
                f"最多支持 {MAX_CONFIG_SHEET_ROWS} 个配置分片",
            )
        )
        return {}, errors
    try:
        header_values = [_text(cell.value) for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    except StopIteration:
        return {}, [_error(1, "表头", f"工作表“{CONFIG_SHEET_TITLE}”缺少表头")]
    if header_values != list(CONFIG_HEADERS):
        return {}, [_error(1, "表头", f"工作表“{CONFIG_SHEET_TITLE}”表头不正确")]

    chunks: dict[str, dict[int, str]] = {}
    first_rows: dict[str, int] = {}
    for row_number, cells in enumerate(sheet.iter_rows(min_row=2, max_col=3), start=2):
        values = [cell.value for cell in cells]
        if all(value is None or _text(value) == "" for value in values):
            continue
        if any(cell.data_type == "f" for cell in cells):
            errors.append(_error(row_number, "配置分片", "不允许使用 Excel 公式，请填写实际值"))
            continue
        reference = _text(values[0])
        if not _CONFIG_REF_RE.fullmatch(reference):
            errors.append(_error(row_number, CONFIG_HEADERS[0], "引用标识格式无效"))
            continue
        chunk_number = _parse_integer(values[1], row_number, CONFIG_HEADERS[1], errors)
        # 分片内容必须保留首尾空白；它可能位于 JSON 字符串值内部，不能复用
        # _text()，否则拼接后会改变配置语义。
        chunk = "" if values[2] is None else str(values[2])
        if chunk_number is None or not chunk:
            if not chunk:
                errors.append(_error(row_number, CONFIG_HEADERS[2], "配置分片不能为空"))
            continue
        reference_chunks = chunks.setdefault(reference, {})
        if chunk_number in reference_chunks:
            errors.append(_error(row_number, CONFIG_HEADERS[1], "同一引用的分片序号重复"))
            continue
        reference_chunks[chunk_number] = chunk
        first_rows.setdefault(reference, row_number)

    resolved: dict[str, str] = {}
    for reference, reference_chunks in chunks.items():
        indexes = sorted(reference_chunks)
        expected = list(range(1, len(indexes) + 1))
        if indexes != expected:
            errors.append(_error(first_rows[reference], CONFIG_HEADERS[1], "配置分片序号必须从 1 连续递增"))
            continue
        resolved[reference] = "".join(reference_chunks[index] for index in indexes)
    return resolved, errors


def _style_data_sheet(sheet) -> None:
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:K{MAX_IMPORT_ROWS + 1}"
    header_fill = PatternFill("solid", fgColor="4F46E5")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    widths = [16, 14, 20, 24, 18, 12, 16, 22, 38, 52, 36]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[chr(64 + index)].width = width
    sheet.row_dimensions[1].height = 32
    _style_data_rows(sheet)
    _apply_data_validation(sheet)


def _write_instructions(sheet, project_id: int | None, project_name: str | None) -> None:
    sheet.sheet_view.showGridLines = False
    sheet.column_dimensions["A"].width = 28
    sheet.column_dimensions["B"].width = 90
    sheet.append(["元素 Excel 填写说明", ""])
    sheet.merge_cells("A1:B1")
    sheet["A1"].fill = PatternFill("solid", fgColor="4F46E5")
    sheet["A1"].font = Font(color="FFFFFF", bold=True, size=14)
    sheet["A1"].alignment = Alignment(horizontal="center")
    target = f"当前项目：{project_name or '项目内导入'}（ID: {project_id or '由接口确定'}）"
    sheet.append(["导入目标", target])
    rows = [
        ["导入规则", "元素ID为空时创建；填写元素ID时更新同项目且由当前用户创建的元素。"],
        ["项目列", "项目ID/项目名称用于导出识别；项目内导入以 URL 中的项目为准，项目ID若填写必须一致。"],
        ["平台", "可填写 android、ios、both；留空默认 both。"],
        ["定位方式", "普通定位填写定位值；smart 定位填写智能定位配置 JSON，定位值留空。"],
        ["智能定位配置", '必须是合法 JSON 对象，例如 {"version":1,"alternatives":[...]}。'],
        ["超长配置", f"导出的超长 smart 配置会放在“{CONFIG_SHEET_TITLE}”工作表中，主表通过引用标识关联，请勿删除或修改分片。"],
        ["事务规则", "文件任意一行失败时整批不写入，请根据错误行修正后重新导入。"],
        ["空白行", "完全空白的数据行会被忽略；最多导入 2,000 条数据。"],
    ]
    for row in rows:
        sheet.append(row)
    for row in sheet.iter_rows(min_row=2, max_col=2):
        row[0].font = Font(bold=True)
        row[0].alignment = Alignment(vertical="top")
        row[1].alignment = Alignment(vertical="top", wrap_text=True)
    sheet.freeze_panes = "A3"


def _new_workbook(project_id: int | None = None, project_name: str | None = None) -> Workbook:
    workbook = Workbook()
    data = workbook.active
    data.title = "元素"
    data.append(list(ELEMENT_HEADERS))
    _style_data_sheet(data)
    instructions = workbook.create_sheet("填写说明")
    _write_instructions(instructions, project_id, project_name)
    return workbook


def _save_workbook(workbook: Workbook) -> bytes:
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def build_template(project_id: int, project_name: str) -> bytes:
    return _save_workbook(_new_workbook(project_id, project_name))


def _safe_text(value: Any) -> str:
    # 让导出的用户字段始终是文本，避免 Excel 将其解释为公式。
    return "" if value is None else str(value)


def build_export(rows: Iterable[tuple[Any, Any]]) -> bytes:
    workbook = _new_workbook()
    sheet = workbook["元素"]
    oversized_configs: dict[str, str] = {}
    for element, project in rows:
        locator_config_text = (
            json.dumps(element.locator_config, ensure_ascii=False, separators=(",", ":"))
            if element.locator_config is not None
            else ""
        )
        if len(locator_config_text) > EXCEL_CELL_MAX_CHARS:
            reference = _config_reference(locator_config_text)
            oversized_configs[reference] = locator_config_text
            locator_config_value = reference
        else:
            locator_config_value = locator_config_text
        values = [
            element.id,
            project.id if project else None,
            _safe_text(project.name if project else None),
            _safe_text(element.name),
            _safe_text(element.page_name),
            _safe_text(element.platform),
            _safe_text(element.scope),
            _safe_text(element.locator_type),
            _safe_text(element.locator_value),
            locator_config_value,
            _safe_text(element.description),
        ]
        sheet.append(values)
        for cell in sheet[sheet.max_row]:
            if cell.column >= 3:
                cell.number_format = "@"
                cell.data_type = "s"
    _style_data_rows(sheet)
    sheet.auto_filter.ref = f"A1:K{max(sheet.max_row, 2)}"
    if oversized_configs:
        _write_config_sheet(workbook, oversized_configs)
    return _save_workbook(workbook)


def parse_import(content: bytes, project_id: int) -> tuple[list[ElementImportRow], list[dict[str, Any]]]:
    _validate_archive(content)
    errors: list[dict[str, Any]] = []
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
    except Exception as exc:
        raise ValueError("无法读取 XLSX 工作簿") from exc
    rows: list[ElementImportRow] = []
    try:
        if "元素" not in workbook.sheetnames:
            return [], [_error(1, "工作表", "必须包含名为“元素”的工作表")]
        sheet = workbook["元素"]
        header_values = [_text(cell.value) for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
        if header_values != list(ELEMENT_HEADERS):
            return [], [_error(1, "表头", "表头与标准模板不一致，请下载最新模板")]
        config_chunks, config_errors = _read_config_sheet(workbook)
        errors.extend(config_errors)

        for row_number, cells in enumerate(sheet.iter_rows(min_row=2), start=2):
            values = [cell.value for cell in cells[: len(ELEMENT_HEADERS)]]
            row_is_blank = all(value is None or _text(value) == "" for value in values)
            if row_is_blank:
                continue
            if row_number > MAX_IMPORT_ROWS + 1:
                errors.append(_error(row_number, "数据行", f"最多导入 {MAX_IMPORT_ROWS} 条数据"))
                break
            if any(cell.data_type == "f" for cell in cells):
                errors.append(_error(row_number, "数据", "不允许使用 Excel 公式，请填写实际值"))
            row_values = dict(zip((ELEMENT_HEADER_FIELDS[h] for h in ELEMENT_HEADERS), values, strict=True))
            element_id = _parse_integer(row_values["element_id"], row_number, "元素ID(element_id)", errors)
            source_project_id = _parse_integer(row_values["project_id"], row_number, "项目ID(project_id)", errors)
            if source_project_id is not None and source_project_id != project_id:
                errors.append(_error(row_number, "项目ID(project_id)", "必须与当前项目一致"))
            name = _text(row_values["name"])
            locator_type = _text(row_values["locator_type"])
            locator_value = _text(row_values["locator_value"]) or None
            config_text = _text(row_values["locator_config"])
            if config_text.startswith(CONFIG_REF_PREFIX):
                if not _CONFIG_REF_RE.fullmatch(config_text):
                    errors.append(_error(row_number, "智能定位配置(locator_config)", "配置引用标识格式无效"))
                    config_text = ""
                else:
                    resolved_config = config_chunks.get(config_text)
                    if resolved_config is None:
                        errors.append(_error(row_number, "智能定位配置(locator_config)", "未找到配置引用对应的分片"))
                        config_text = ""
                    else:
                        config_text = resolved_config
            locator_config: dict[str, Any] | None = None
            if config_text:
                try:
                    parsed = json.loads(config_text)
                    if not isinstance(parsed, dict):
                        raise ValueError
                    locator_config = parsed
                except (TypeError, ValueError, json.JSONDecodeError):
                    errors.append(_error(row_number, "智能定位配置(locator_config)", "必须是合法 JSON 对象"))
            rows.append(
                ElementImportRow(
                    row_number=row_number,
                    element_id=element_id,
                    project_id=source_project_id,
                    data={
                        "name": name,
                        "page_name": _text(row_values["page_name"]) or None,
                        "platform": _text(row_values["platform"]) or "both",
                        "scope": _text(row_values["scope"]) or "all",
                        "locator_type": locator_type,
                        "locator_value": locator_value,
                        "locator_config": locator_config,
                        "description": _text(row_values["description"]) or None,
                    },
                )
            )
        return rows, errors
    finally:
        workbook.close()
