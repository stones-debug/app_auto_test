import json
from collections.abc import Sequence
from hashlib import sha256
from io import BytesIO
from types import SimpleNamespace
from typing import Any

import pytest
from httpx import ASGITransport, AsyncClient
from openpyxl import Workbook, load_workbook
from sqlalchemy import update

from app.core.database import SessionLocal
from app.main import app
from app.models import TestElement as ElementModel
from app.services.element_excel import (
    CONFIG_CHUNK_SIZE,
    CONFIG_HEADERS,
    CONFIG_REF_PREFIX,
    CONFIG_SHEET_TITLE,
    ELEMENT_HEADERS,
    EXCEL_CELL_MAX_CHARS,
    build_export,
    parse_import,
)

OWNER = {"username": "pytest_owner", "email": "owner@tl-tek.com", "password": "test123"}


@pytest.fixture
async def client():
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as c:
        yield c


async def _setup(client: AsyncClient) -> tuple[dict, int]:
    """注册 owner 并创建项目，返回 (headers, project_id)。"""
    reg = await client.post("/api/auth/register", json=OWNER)
    assert reg.status_code == 201
    token = reg.json()["access_token"]
    headers = {"Authorization": f"Bearer {token}"}
    created = await client.post(
        "/api/projects", json={"name": "元素项目"}, headers=headers
    )
    assert created.status_code == 201
    return headers, created.json()["id"]


async def test_module_crud(client: AsyncClient):
    headers, project_id = await _setup(client)

    created = await client.post(
        f"/api/projects/{project_id}/modules",
        json={"name": "登录模块"},
        headers=headers,
    )
    assert created.status_code == 201
    module_id = created.json()["id"]

    # 子模块
    child = await client.post(
        f"/api/projects/{project_id}/modules",
        json={"name": "子模块", "parent_id": module_id},
        headers=headers,
    )
    assert child.status_code == 201

    listing = await client.get(f"/api/projects/{project_id}/modules", headers=headers)
    assert listing.status_code == 200
    names = [m["name"] for m in listing.json()]
    assert "登录模块" in names

    # 子模块通过 parent_id 过滤
    child_list = await client.get(
        f"/api/projects/{project_id}/modules?parent_id={module_id}", headers=headers
    )
    assert len(child_list.json()) == 1
    assert child_list.json()[0]["name"] == "子模块"

    updated = await client.put(
        f"/api/modules/{module_id}", json={"name": "改名模块"}, headers=headers
    )
    assert updated.status_code == 200
    assert updated.json()["name"] == "改名模块"

    deleted = await client.delete(f"/api/modules/{module_id}", headers=headers)
    assert deleted.status_code == 204


async def test_element_crud_and_usage(client: AsyncClient):
    headers, project_id = await _setup(client)

    created = await client.post(
        f"/api/projects/{project_id}/elements",
        json={
            "name": "登录按钮",
            "page_name": "登录页",
            "platform": "both",
            "locator_type": "resource_id",
            "locator_value": "btn_login",
            "description": "登录按钮",
        },
        headers=headers,
    )
    assert created.status_code == 201
    element_id = created.json()["id"]

    listing = await client.get(
        f"/api/projects/{project_id}/elements?keyword=登录", headers=headers
    )
    assert listing.status_code == 200
    assert listing.json()["total"] == 1
    assert listing.json()["items"][0]["locator_type"] == "resource_id"
    assert listing.json()["items"][0]["locator_value"] == "btn_login"

    detail = await client.get(f"/api/elements/{element_id}", headers=headers)
    assert detail.status_code == 200

    updated = await client.put(
        f"/api/elements/{element_id}",
        json={"locator_value": "btn_login_new"},
        headers=headers,
    )
    assert updated.status_code == 200
    assert updated.json()["locator_value"] == "btn_login_new"

    # usage：无用例引用时为空
    usage = await client.get(f"/api/elements/{element_id}/usage", headers=headers)
    assert usage.status_code == 200
    assert usage.json() == []

    deleted = await client.delete(f"/api/elements/{element_id}", headers=headers)
    assert deleted.status_code == 204


async def test_element_pages_grouping(client: AsyncClient):
    """B3：element-pages 按 page_name 分组统计，NULL 归为"未分组"。"""
    headers, project_id = await _setup(client)
    for name, page in [
        ("元素A", "登录页"),
        ("元素B", "登录页"),
        ("元素C", None),
        ("元素D", "首页"),
    ]:
        resp = await client.post(
            f"/api/projects/{project_id}/elements",
            json={"name": name, "page_name": page, "locator_type": "id", "locator_value": name},
            headers=headers,
        )
        assert resp.status_code == 201

    pages = (
        await client.get(f"/api/projects/{project_id}/element-pages", headers=headers)
    ).json()
    counts = {p["page_name"]: p["count"] for p in pages}
    assert counts["登录页"] == 2
    assert counts["未分组"] == 1
    assert counts["首页"] == 1


async def test_element_usage_step_orders(client: AsyncClient):
    """B3：usage 返回引用该元素的步骤 step_orders（仅 steps）。"""
    headers, project_id = await _setup(client)
    el = await client.post(
        f"/api/projects/{project_id}/elements",
        json={"name": "用户名", "locator_type": "id", "locator_value": "username"},
        headers=headers,
    )
    element_id = el.json()["id"]
    # 两个步骤引用该元素（order 1/2），一个断言也引用（不计入 step_orders）
    case = await client.post(
        f"/api/projects/{project_id}/cases",
        json={
            "name": "引用用例",
            "steps": [
                {"order": 1, "action": "input", "element_id": element_id, "params": {"value": "admin"}},
                {"order": 2, "action": "click", "element_id": element_id, "params": {}},
            ],
            "assertions": [
                {"order": 1, "type": "text_equals", "element_id": element_id, "params": {"expected": "x"}}
            ],
        },
        headers=headers,
    )
    assert case.status_code == 201

    usage = (await client.get(f"/api/elements/{element_id}/usage", headers=headers)).json()
    assert len(usage) == 1
    assert usage[0]["case_name"] == "引用用例"
    assert usage[0]["step_orders"] == [1, 2]


async def test_element_permission(client: AsyncClient):
    """非成员不能创建元素。"""
    reg = await client.post("/api/auth/register", json=OWNER)
    token = reg.json()["access_token"]
    headers = {"Authorization": f"Bearer {token}"}
    project_id = (
        await client.post("/api/projects", json={"name": "权限项目"}, headers=headers)
    ).json()["id"]

    # 另一用户
    intruder = {"username": "pytest_intruder", "email": "intruder@tl-tek.com", "password": "test123"}
    intruder_token = (await client.post("/api/auth/register", json=intruder)).json()["access_token"]
    intruder_headers = {"Authorization": f"Bearer {intruder_token}"}

    resp = await client.post(
        f"/api/projects/{project_id}/elements",
        json={
            "name": "越权元素",
            "locator_type": "id",
            "locator_value": "x",
        },
        headers=intruder_headers,
    )
    assert resp.status_code == 403


# ---------- CR-18：模块父级归属与循环 ----------


async def test_module_update_rejects_cycle(client: AsyncClient):
    headers, project_id = await _setup(client)
    a = (await client.post(
        f"/api/projects/{project_id}/modules", json={"name": "A"}, headers=headers
    )).json()["id"]
    b = (await client.post(
        f"/api/projects/{project_id}/modules", json={"name": "B", "parent_id": a}, headers=headers
    )).json()["id"]
    c = (await client.post(
        f"/api/projects/{project_id}/modules", json={"name": "C", "parent_id": b}, headers=headers
    )).json()["id"]

    # 自身为父 → 400
    self_parent = await client.put(
        f"/api/modules/{a}", json={"parent_id": a}, headers=headers
    )
    assert self_parent.status_code == 400

    # a 的父设为后代 c → 循环 → 400
    cycle = await client.put(
        f"/api/modules/{a}", json={"parent_id": c}, headers=headers
    )
    assert cycle.status_code == 400
    assert "循环" in cycle.json()["detail"]

    # 跨项目父 → 404（已存在校验）
    p2 = (await client.post("/api/projects", json={"name": "项目2"}, headers=headers)).json()["id"]
    cross = await client.put(
        f"/api/modules/{a}", json={"parent_id": p2}, headers=headers
    )
    assert cross.status_code == 404


async def test_module_update_can_clear_parent(client: AsyncClient):
    """CR-25：显式传 parent_id=0 清空父模块。"""
    headers, project_id = await _setup(client)
    a = (await client.post(
        f"/api/projects/{project_id}/modules", json={"name": "A"}, headers=headers
    )).json()["id"]
    b = (await client.post(
        f"/api/projects/{project_id}/modules", json={"name": "B", "parent_id": a}, headers=headers
    )).json()["id"]

    cleared = await client.put(f"/api/modules/{b}", json={"parent_id": 0}, headers=headers)
    assert cleared.status_code == 200
    assert cleared.json()["parent_id"] is None


async def test_element_list_page_name_and_locator_filters(client: AsyncClient):
    """F8：元素列表支持 page_name（含未分组）与 locator_type 过滤。"""
    headers, project_id = await _setup(client)
    for name, page, locator in [
        ("元素1", "登录页", "id"),
        ("元素2", "登录页", "xpath"),
        ("元素3", None, "id"),
    ]:
        await client.post(
            f"/api/projects/{project_id}/elements",
            json={"name": name, "page_name": page, "locator_type": locator, "locator_value": name},
            headers=headers,
        )

    by_page = (await client.get(
        f"/api/projects/{project_id}/elements?page_name=登录页", headers=headers
    )).json()
    assert by_page["total"] == 2

    by_unset = (await client.get(
        f"/api/projects/{project_id}/elements?page_name=未分组", headers=headers
    )).json()
    assert by_unset["total"] == 1

    by_locator = (await client.get(
        f"/api/projects/{project_id}/elements?locator_type=xpath", headers=headers
    )).json()
    assert by_locator["total"] == 1
    assert by_locator["items"][0]["name"] == "元素2"


# ---------- V3：元素库全局化 / 仅创建者可改 / 复制 / 自定义分组 ----------


async def _register_user(client: AsyncClient, username: str, email: str) -> dict:
    token = (await client.post(
        "/api/auth/register",
        json={"username": username, "email": email, "password": "test123"},
    )).json()["access_token"]
    return {"Authorization": f"Bearer {token}"}


async def test_element_global_list_and_project_filter(client: AsyncClient):
    """V3：全库列表可见两个项目的元素且带项目名；可按 project_id 筛选。"""
    h1, p1 = await _setup(client)
    p2 = (await client.post("/api/projects", json={"name": "元素项目2"}, headers=h1)).json()["id"]

    e1 = (await client.post(
        "/api/elements",
        headers=h1,
        json={"project_id": p1, "name": "全库元素A", "locator_type": "id", "locator_value": "a"},
    )).json()
    e2 = (await client.post(
        "/api/elements",
        headers=h1,
        json={"project_id": p2, "name": "全库元素B", "locator_type": "id", "locator_value": "b"},
    )).json()
    assert e1["project_name"] == "元素项目"
    assert e2["project_name"] == "元素项目2"

    listing = (await client.get("/api/elements", headers=h1)).json()
    names = {i["name"]: i for i in listing["items"]}
    assert "全库元素A" in names and "全库元素B" in names
    assert names["全库元素A"]["project_name"] == "元素项目"
    assert names["全库元素A"]["created_by_name"] == "pytest_owner"

    filtered = (await client.get(f"/api/elements?project_id={p2}", headers=h1)).json()
    assert filtered["total"] == 1
    assert filtered["items"][0]["name"] == "全库元素B"


async def test_element_pages_filter_groups_by_project(client: AsyncClient):
    """项目元素库只返回项目实际使用的分组，并保留其层级父节点。"""
    headers, project_id = await _setup(client)
    other_project_id = (await client.post(
        "/api/projects", json={"name": "元素分组过滤项目2"}, headers=headers
    )).json()["id"]

    root = (await client.post(
        "/api/elements/groups", headers=headers, json={"name": "过滤项目一父页面"}
    )).json()
    child = (await client.post(
        "/api/elements/groups",
        headers=headers,
        json={"name": "过滤项目一子页面", "parent_id": root["id"]},
    )).json()
    other_group = (await client.post(
        "/api/elements/groups", headers=headers, json={"name": "过滤项目二页面"}
    )).json()
    empty_group = (await client.post(
        "/api/elements/groups", headers=headers, json={"name": "过滤项目空页面"}
    )).json()

    for project, page_name, element_name in (
        (project_id, child["name"], "过滤项目一元素"),
        (other_project_id, other_group["name"], "过滤项目二元素"),
    ):
        response = await client.post(
            "/api/elements",
            headers=headers,
            json={
                "project_id": project,
                "name": element_name,
                "page_name": page_name,
                "locator_type": "id",
                "locator_value": element_name,
            },
        )
        assert response.status_code == 201

    project_pages = (await client.get(
        f"/api/elements/pages?project_id={project_id}", headers=headers
    )).json()
    project_names = {item["page_name"] for item in project_pages}
    assert {root["name"], child["name"]} <= project_names
    assert "过滤项目二页面" not in project_names
    assert "过滤项目空页面" not in project_names

    other_pages = (await client.get(
        f"/api/elements/pages?project_id={other_project_id}", headers=headers
    )).json()
    other_names = {item["page_name"] for item in other_pages}
    assert other_group["name"] in other_names
    assert child["name"] not in other_names

    # 分组是全局资源且 created_by 关联用户，测试结束前显式清理，避免影响后续用户清理夹具。
    for group_id in (child["id"], root["id"], other_group["id"], empty_group["id"]):
        assert (
            await client.delete(f"/api/elements/groups/{group_id}", headers=headers)
        ).status_code == 204


async def test_element_creator_only_edit_delete(client: AsyncClient):
    """V3：非创建者（管理员也不行）不能编辑/删除元素；创建者可以。"""
    h_owner, p1 = await _setup(client)
    h_admin = await _register_user(client, "pytest_admin", "admin2@tl-tek.com")

    # admin 加入项目（owner 邀请）
    me = (await client.get("/api/auth/me", headers=h_admin)).json()
    await client.post(f"/api/projects/{p1}/members", headers=h_owner, json={"user_id": me["id"], "role": "admin"})

    el = (await client.post(
        "/api/elements",
        headers=h_owner,
        json={"project_id": p1, "name": "创建者元素", "locator_type": "id", "locator_value": "c"},
    )).json()
    el_id = el["id"]

    # 非创建者 admin 更新 → 403
    denied = await client.put(
        f"/api/elements/{el_id}", json={"locator_value": "hack"}, headers=h_admin
    )
    assert denied.status_code == 403
    assert "创建者" in denied.json()["detail"]
    # 非创建者删除 → 403
    assert (await client.delete(f"/api/elements/{el_id}", headers=h_admin)).status_code == 403

    # 创建者更新 → 200
    ok = (await client.put(
        f"/api/elements/{el_id}", json={"locator_value": "creator_update"}, headers=h_owner
    )).json()
    assert ok["locator_value"] == "creator_update"
    # 创建者删除 → 204
    assert (await client.delete(f"/api/elements/{el_id}", headers=h_owner)).status_code == 204


async def test_element_copy_creates_for_current_user(client: AsyncClient):
    """V3：复制按钮按当前用户新建相同元素（归属原项目）。"""
    h_owner, p1 = await _setup(client)
    h_other = await _register_user(client, "pytest_other", "other@tl-tek.com")
    el = (await client.post(
        "/api/elements",
        headers=h_owner,
        json={"project_id": p1, "name": "被复制", "page_name": "登录页", "locator_type": "id", "locator_value": "c1"},
    )).json()

    copied = (await client.post(f"/api/elements/{el['id']}/copy", headers=h_other)).json()
    assert copied["id"] != el["id"]
    assert copied["name"] == "被复制"
    assert copied["page_name"] == "登录页"
    assert copied["locator_value"] == "c1"
    assert copied["created_by_name"] == "pytest_other"
    assert copied["project_id"] == p1


async def test_element_update_can_move_project(client: AsyncClient):
    """V3：编辑时可迁移元素项目（创建者+目标项目写权限）；无权限目标项目 403；非创建者 403。"""
    h_owner, p1 = await _setup(client)
    p2 = (await client.post("/api/projects", json={"name": "迁移目标项目"}, headers=h_owner)).json()["id"]
    h_stranger = await _register_user(client, "pytest_stranger", "stranger@tl-tek.com")
    p3 = (await client.post("/api/projects", json={"name": "他人项目"}, headers=h_stranger)).json()["id"]

    el = (await client.post(
        "/api/elements",
        headers=h_owner,
        json={"project_id": p1, "name": "待迁移", "locator_type": "id", "locator_value": "mv"},
    )).json()
    el_id = el["id"]

    # 无权限目标项目（他人项目）→ 403
    denied = await client.put(
        f"/api/elements/{el_id}", json={"project_id": p3}, headers=h_owner
    )
    assert denied.status_code == 403

    # 迁移成功
    moved = (await client.put(
        f"/api/elements/{el_id}", json={"project_id": p2}, headers=h_owner
    )).json()
    assert moved["project_id"] == p2
    assert moved["project_name"] == "迁移目标项目"
    assert moved["created_by_name"] == "pytest_owner"

    # 旧项目列表不含、新项目包含
    old = (await client.get(f"/api/elements?project_id={p1}", headers=h_owner)).json()
    assert all(i["id"] != el_id for i in old["items"])
    new = (await client.get(f"/api/elements?project_id={p2}", headers=h_owner)).json()
    assert any(i["id"] == el_id for i in new["items"])


async def test_element_group_custom(client: AsyncClient):
    """自定义分组支持创建、重命名和删除后自动转为未分组。"""
    h1, project_id = await _setup(client)
    g = (await client.post("/api/elements/groups", headers=h1, json={"name": "我的新分组"})).json()
    group_id = g["id"]
    assert g["name"] == "我的新分组"

    conflict = await client.post("/api/elements/groups", headers=h1, json={"name": "我的新分组"})
    assert conflict.status_code == 409

    child = await client.post(
        "/api/elements/groups",
        headers=h1,
        json={"name": "子页面", "parent_id": group_id},
    )
    assert child.status_code == 201
    child_id = child.json()["id"]
    assert child.json()["parent_id"] == group_id

    renamed = await client.put(
        f"/api/elements/groups/{group_id}",
        headers=h1,
        json={"name": "重命名分组"},
    )
    assert renamed.status_code == 200
    assert renamed.json()["name"] == "重命名分组"

    duplicate = await client.post(
        "/api/elements/groups", headers=h1, json={"name": "已有分组"}
    )
    assert duplicate.status_code == 201
    rename_conflict = await client.put(
        f"/api/elements/groups/{group_id}",
        headers=h1,
        json={"name": "已有分组"},
    )
    assert rename_conflict.status_code == 409

    element = await client.post(
        "/api/elements",
        headers=h1,
        json={
            "project_id": project_id,
            "name": "分组内元素",
            "page_name": "重命名分组",
            "locator_type": "id",
            "locator_value": "grouped",
        },
    )
    assert element.status_code == 201
    element_id = element.json()["id"]

    renamed_element = await client.get(f"/api/elements/{element_id}", headers=h1)
    assert renamed_element.status_code == 200
    assert renamed_element.json()["page_name"] == "重命名分组"

    pages = (await client.get("/api/elements/pages", headers=h1)).json()
    names = {p["page_name"]: p["count"] for p in pages}
    assert names["重命名分组"] == 1

    h_other = await _register_user(client, "pytest_group2", "group2@tl-tek.com")
    other_renamed = await client.put(
        f"/api/elements/groups/{group_id}",
        headers=h_other,
        json={"name": "他人重命名"},
    )
    assert other_renamed.status_code == 200
    assert other_renamed.json()["name"] == "他人重命名"
    assert (await client.delete(f"/api/elements/groups/{group_id}", headers=h_other)).status_code == 204
    ungrouped = await client.get(f"/api/elements/{element_id}", headers=h1)
    assert ungrouped.status_code == 200
    assert ungrouped.json()["page_name"] is None
    pages_after = (await client.get("/api/elements/pages", headers=h1)).json()
    assert "他人重命名" not in {p["page_name"] for p in pages_after}
    child_page = next(p for p in pages_after if p["page_name"] == "子页面")
    assert child_page["parent_id"] is None
    assert (
        await client.delete(f"/api/elements/groups/{duplicate.json()['id']}", headers=h1)
    ).status_code == 204
    assert (
        await client.delete(f"/api/elements/groups/{child_id}", headers=h1)
    ).status_code == 204


async def test_element_page_group_reserved_names_are_rejected(client: AsyncClient):
    """页面分组保留名不能创建，也不能通过元素页面名称间接创建。"""
    headers, project_id = await _setup(client)
    for name in ("all", "ALL", "全部", "未分组", "  all  "):
        response = await client.post(
            "/api/elements/groups", headers=headers, json={"name": name}
        )
        assert response.status_code == 422

    element = await client.post(
        "/api/elements",
        headers=headers,
        json={
            "project_id": project_id,
            "name": "保留名测试元素",
            "page_name": "all",
            "locator_type": "id",
            "locator_value": "reserved",
        },
    )
    assert element.status_code == 422

    group = (
        await client.post("/api/elements/groups", headers=headers, json={"name": "普通分组"})
    ).json()
    for name in ("全部", "未分组", "ALL"):
        response = await client.put(
            f"/api/elements/groups/{group['id']}",
            headers=headers,
            json={"name": name},
        )
        assert response.status_code == 422
    assert (
        await client.delete(f"/api/elements/groups/{group['id']}", headers=headers)
    ).status_code == 204


async def test_global_ungrouped_handles_null_blank_and_whitespace(client: AsyncClient):
    """未分组统计和列表统一覆盖 NULL、空字符串、纯空格及清空分组。"""
    headers, project_id = await _setup(client)
    created = []
    for name, page_name in [
        ("NULL元素", None),
        ("历史空串元素", "临时页"),
        ("历史空格元素", "临时页"),
    ]:
        response = await client.post(
            "/api/elements",
            headers=headers,
            json={
                "project_id": project_id,
                "name": name,
                "page_name": page_name,
                "locator_type": "id",
                "locator_value": name,
            },
        )
        assert response.status_code == 201
        created.append(response.json())

    # 模拟修复前已落库的空字符串/纯空格历史数据。
    async with SessionLocal() as db:
        await db.execute(
            update(ElementModel)
            .where(ElementModel.id == created[1]["id"])
            .values(page_name="")
        )
        await db.execute(
            update(ElementModel)
            .where(ElementModel.id == created[2]["id"])
            .values(page_name="   ")
        )
        await db.commit()

    named = (
        await client.post(
            "/api/elements",
            headers=headers,
            json={
                "project_id": project_id,
                "name": "可清空元素",
                "page_name": "  登录页  ",
                "locator_type": "id",
                "locator_value": "clearable",
            },
        )
    ).json()
    assert named["page_name"] == "登录页"
    cleared = await client.put(
        f"/api/elements/{named['id']}",
        headers=headers,
        json={"page_name": "   "},
    )
    assert cleared.status_code == 200
    assert cleared.json()["page_name"] is None

    pages = (await client.get("/api/elements/pages", headers=headers)).json()
    counts = {item["page_name"]: item["count"] for item in pages}
    assert counts["未分组"] == 4
    assert "" not in counts

    ungrouped = (
        await client.get("/api/elements?page_name=未分组", headers=headers)
    ).json()
    assert ungrouped["total"] == 4
    assert {item["name"] for item in ungrouped["items"]} == {
        "NULL元素",
        "历史空串元素",
        "历史空格元素",
        "可清空元素",
    }


async def test_element_scope_default_and_edit(client: AsyncClient):
    """适用范围：不填默认 all；创建可填自定义；编辑可修改；空白归为 all。"""
    headers, project_id = await _setup(client)

    # 不填 scope → 默认 all
    default_el = await client.post(
        f"/api/projects/{project_id}/elements",
        json={
            "name": "默认范围元素",
            "locator_type": "id",
            "locator_value": "default_scope",
        },
        headers=headers,
    )
    assert default_el.status_code == 201
    assert default_el.json()["scope"] == "all"

    # 显式填自定义范围
    scoped = await client.post(
        f"/api/projects/{project_id}/elements",
        json={
            "name": "DVR范围元素",
            "locator_type": "id",
            "locator_value": "dvr_scope",
            "scope": "DVR",
        },
        headers=headers,
    )
    assert scoped.status_code == 201
    assert scoped.json()["scope"] == "DVR"
    scoped_id = scoped.json()["id"]

    # 编辑修改适用范围
    updated = await client.put(
        f"/api/elements/{scoped_id}",
        json={"scope": "网约车"},
        headers=headers,
    )
    assert updated.status_code == 200
    assert updated.json()["scope"] == "网约车"

    # 编辑空白 → 兜底 all
    cleared = await client.put(
        f"/api/elements/{scoped_id}",
        json={"scope": "   "},
        headers=headers,
    )
    assert cleared.status_code == 200
    assert cleared.json()["scope"] == "all"

    # 列表与详情均返回 scope
    listing = await client.get(
        f"/api/projects/{project_id}/elements?keyword=默认范围元素", headers=headers
    )
    assert listing.json()["total"] == 1
    assert listing.json()["items"][0]["scope"] == "all"

    detail = await client.get(f"/api/elements/{default_el.json()['id']}", headers=headers)
    assert detail.json()["scope"] == "all"


# ---------- 智能元素定位（smart locator） ----------

SMART_CONFIG = {
    "version": 1,
    "alternatives": [
        {
            "anchor": [{"attribute": "text", "operator": "equals", "value": "登录"}],
            "path": [{"axis": "parent", "depth": 1}],
            "target": [{"attribute": "resource_id", "operator": "contains", "value": "btn"}],
        }
    ],
    "search": {"scroll": True, "direction": "up", "max_swipes": 8, "duration_ms": 500, "settle_ms": 300},
    # model_dump 会显式序列化默认 None 字段（index），此处与其保持一致
    "selection": {"policy": "unique", "index": None},
}


async def test_element_create_smart(client: AsyncClient):
    """创建 smart 元素成功：合法 locator_config，locator_value 不传/null，返回 config 完整。"""
    headers, project_id = await _setup(client)
    created = await client.post(
        "/api/elements",
        headers=headers,
        json={"project_id": project_id, "name": "智能定位按钮", "locator_type": "smart", "locator_config": SMART_CONFIG},
    )
    assert created.status_code == 201
    data = created.json()
    assert data["locator_type"] == "smart"
    assert data["locator_value"] is None
    assert data["locator_config"] == SMART_CONFIG

    # 显式传 locator_value=null 同样合法
    created2 = await client.post(
        "/api/elements",
        headers=headers,
        json={"project_id": project_id, "name": "智能定位按钮2", "locator_type": "smart", "locator_value": None, "locator_config": SMART_CONFIG},
    )
    assert created2.status_code == 201
    assert created2.json()["locator_value"] is None


async def test_element_create_smart_validation(client: AsyncClient):
    """smart 缺 locator_config → 422；smart 传非空 locator_value → 422。"""
    headers, project_id = await _setup(client)
    missing = await client.post(
        "/api/elements",
        headers=headers,
        json={"project_id": project_id, "name": "缺配置", "locator_type": "smart"},
    )
    assert missing.status_code == 422

    with_value = await client.post(
        "/api/elements",
        headers=headers,
        json={
            "project_id": project_id,
            "name": "带值",
            "locator_type": "smart",
            "locator_value": "x",
            "locator_config": SMART_CONFIG,
        },
    )
    assert with_value.status_code == 422


async def test_element_create_smart_combo_positive_cases(client: AsyncClient):
    """合法组合不被误杀：纯 regex target（Strategy U）、regex+equals（均可 UiAutomator）、
    anchor-only（无 path）、anchor+path 含 contains（XPath 可表达）。"""
    headers, project_id = await _setup(client)

    def _body(config):
        return {"project_id": project_id, "name": "合法组合", "locator_type": "smart", "locator_config": config}

    positive = [
        # 纯 regex（Strategy U）
        {"version": 1, "alternatives": [{"target": [{"attribute": "text", "operator": "regex", "value": "\\d+"}]}]},
        # regex + equals（均可 UiAutomator 表达）
        {"version": 1, "alternatives": [{"target": [
            {"attribute": "text", "operator": "regex", "value": "^\\d+$"},
            {"attribute": "text", "operator": "equals", "value": "登录"},
        ]}]},
        # anchor-only（无 path）
        {"version": 1, "alternatives": [{"anchor": [{"attribute": "text", "operator": "equals", "value": "登录"}], "target": [{"attribute": "resource_id", "operator": "contains", "value": "btn"}]}]},
        # anchor+path 含 contains target（XPath 表达，无 regex）
        SMART_CONFIG,
    ]
    for config in positive:
        resp = await client.post("/api/elements", headers=headers, json=_body(config))
        assert resp.status_code == 201, f"config {config} should be accepted"


async def test_element_create_smart_invalid_configs(client: AsyncClient):
    """非法 config 系列全部 422：未知 attribute/operator、布尔属性非 equals、无效 regex、超限结构。"""
    headers, project_id = await _setup(client)

    def _body(config):
        return {"project_id": project_id, "name": "非法", "locator_type": "smart", "locator_config": config}

    bad_configs = [
        # 未知 attribute
        {"version": 1, "alternatives": [{"target": [{"attribute": "foo", "operator": "equals", "value": "x"}]}]},
        # 未知 operator
        {"version": 1, "alternatives": [{"target": [{"attribute": "text", "operator": "frobnicate", "value": "x"}]}]},
        # 布尔属性用 contains
        {"version": 1, "alternatives": [{"target": [{"attribute": "clickable", "operator": "contains", "value": True}]}]},
        # 无效正则 "["
        {"version": 1, "alternatives": [{"target": [{"attribute": "text", "operator": "regex", "value": "["}]}]},
        # path 4 段（上限 3）
        {
            "version": 1,
            "alternatives": [
                {
                    "target": [{"attribute": "text", "operator": "equals", "value": "x"}],
                    "path": [
                        {"axis": "parent", "depth": 1},
                        {"axis": "parent", "depth": 1},
                        {"axis": "parent", "depth": 1},
                        {"axis": "parent", "depth": 1},
                    ],
                }
            ],
        },
        # ancestor depth 6（上限 5）
        {
            "version": 1,
            "alternatives": [
                {
                    "target": [{"attribute": "text", "operator": "equals", "value": "x"}],
                    "path": [{"axis": "ancestor", "depth": 6}],
                }
            ],
        },
        # alternatives 11 个（上限 10）
        {
            "version": 1,
            "alternatives": [
                {"target": [{"attribute": "text", "operator": "equals", "value": f"x{i}"}]} for i in range(11)
            ],
        },
        # max_swipes 30（上限 20）
        {
            "version": 1,
            "alternatives": [{"target": [{"attribute": "text", "operator": "equals", "value": "x"}]}],
            "search": {"max_swipes": 30},
        },
        # path 无 anchor（组合规则）
        {"version": 1, "alternatives": [{"target": [{"attribute": "text", "operator": "equals", "value": "x"}], "path": [{"axis": "child"}]}]},
        # anchor/path 相对定位禁用 regex
        {"version": 1, "alternatives": [{"anchor": [{"attribute": "text", "operator": "equals", "value": "锚点"}], "path": [{"axis": "child"}], "target": [{"attribute": "text", "operator": "regex", "value": "\\d+"}]}]},
        # regex 与需 XPath 表达的条件混用
        {"version": 1, "alternatives": [{"target": [{"attribute": "text", "operator": "regex", "value": "\\d+"}, {"attribute": "text", "operator": "ends_with", "value": "页"}]}]},
    ]
    for config in bad_configs:
        resp = await client.post("/api/elements", headers=headers, json=_body(config))
        assert resp.status_code == 422, f"config {config} should be rejected"


async def test_element_create_ordinary_still_requires_value(client: AsyncClient):
    """普通元素（id）创建仍必须非空 locator_value，传 locator_config → 422。"""
    headers, project_id = await _setup(client)
    with_config = await client.post(
        "/api/elements",
        headers=headers,
        json={
            "project_id": project_id,
            "name": "普通带配置",
            "locator_type": "id",
            "locator_value": "x",
            "locator_config": SMART_CONFIG,
        },
    )
    assert with_config.status_code == 422

    missing = await client.post(
        "/api/elements",
        headers=headers,
        json={"project_id": project_id, "name": "普通缺值", "locator_type": "id"},
    )
    assert missing.status_code == 422


async def test_element_update_smart_roundtrip(client: AsyncClient):
    """更新元素 id→smart（携 config）成功；smart→id（携 value）成功。"""
    headers, project_id = await _setup(client)
    el = (
        await client.post(
            "/api/elements",
            headers=headers,
            json={"project_id": project_id, "name": "可切换", "locator_type": "id", "locator_value": "old_id"},
        )
    ).json()
    el_id = el["id"]

    to_smart = await client.put(
        f"/api/elements/{el_id}",
        headers=headers,
        json={"locator_type": "smart", "locator_config": SMART_CONFIG},
    )
    assert to_smart.status_code == 200
    assert to_smart.json()["locator_type"] == "smart"
    assert to_smart.json()["locator_config"] == SMART_CONFIG
    # 归一化：改 smart 后旧 locator_value 自动清空为 null（满足 CHECK 约束、避免回显残留）
    assert to_smart.json()["locator_value"] is None

    to_id = await client.put(
        f"/api/elements/{el_id}",
        headers=headers,
        json={"locator_type": "id", "locator_value": "new_id"},
    )
    assert to_id.status_code == 200
    assert to_id.json()["locator_type"] == "id"
    assert to_id.json()["locator_value"] == "new_id"
    # 归一化：改回普通后旧 locator_config 自动清空为 null
    assert to_id.json()["locator_config"] is None


async def test_element_copy_smart_preserves_config(client: AsyncClient):
    """复制 smart 元素保留 locator_config。"""
    h_owner, p1 = await _setup(client)
    el = (
        await client.post(
            "/api/elements",
            headers=h_owner,
            json={"project_id": p1, "name": "被复制智能", "locator_type": "smart", "locator_config": SMART_CONFIG},
        )
    ).json()
    copied = (await client.post(f"/api/elements/{el['id']}/copy", headers=h_owner)).json()
    assert copied["id"] != el["id"]
    assert copied["locator_type"] == "smart"
    assert copied["locator_config"] == SMART_CONFIG


async def test_element_list_detail_returns_locator_config(client: AsyncClient):
    """列表/详情均返回 locator_config。"""
    headers, project_id = await _setup(client)
    el = (
        await client.post(
            "/api/elements",
            headers=headers,
            json={"project_id": project_id, "name": "列表智能", "locator_type": "smart", "locator_config": SMART_CONFIG},
        )
    ).json()
    el_id = el["id"]

    listing = (await client.get("/api/elements?keyword=列表智能", headers=headers)).json()
    assert listing["total"] == 1
    assert listing["items"][0]["locator_config"] == SMART_CONFIG

    detail = (await client.get(f"/api/elements/{el_id}", headers=headers)).json()
    assert detail["locator_config"] == SMART_CONFIG


def _element_xlsx(rows: Sequence[Sequence[Any]]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "元素"
    sheet.append(list(ELEMENT_HEADERS))
    for row in rows:
        sheet.append(list(row))
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


async def test_element_excel_template_import_export_and_update(client: AsyncClient):
    """模板、普通/智能元素导入、全量导出及按 ID 更新。"""
    headers, project_id = await _setup(client)
    template = await client.get(f"/api/projects/{project_id}/elements/import-template", headers=headers)
    assert template.status_code == 200
    template_book = load_workbook(BytesIO(template.content), read_only=True)
    assert template_book.sheetnames == ["元素", "填写说明"]
    assert [cell.value for cell in next(template_book["元素"].iter_rows(max_row=1))] == list(ELEMENT_HEADERS)
    template_book.close()

    smart_json = json.dumps(SMART_CONFIG, ensure_ascii=False, separators=(",", ":"))
    imported = await client.post(
        f"/api/projects/{project_id}/elements/import",
        headers=headers,
        files={
            "file": (
                "elements.xlsx",
                _element_xlsx([
                    [None, project_id, "元素项目", "Excel按钮", "首页", "android", "all", "resource_id", "btn_excel", None, "批量导入"],
                    [None, project_id, "元素项目", "Excel智能", "首页", "android", "all", "smart", None, smart_json, "智能"],
                ]),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert imported.status_code == 200, imported.text
    assert imported.json() == {"created": 2, "updated": 0, "total": 2}

    exported = await client.get(f"/api/elements/export?project_id={project_id}&keyword=Excel", headers=headers)
    assert exported.status_code == 200
    exported_book = load_workbook(BytesIO(exported.content), read_only=True, data_only=False)
    exported_rows = list(exported_book["元素"].iter_rows(min_row=2, values_only=True))
    assert len(exported_rows) == 2
    button_row = next(row for row in exported_rows if row[3] == "Excel按钮")
    exported_book.close()

    updated = list(button_row)
    updated[3] = "Excel按钮-已更新"
    update_response = await client.post(
        f"/api/projects/{project_id}/elements/import",
        headers=headers,
        files={"file": ("update.xlsx", _element_xlsx([updated]), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert update_response.status_code == 200, update_response.text
    assert update_response.json() == {"created": 0, "updated": 1, "total": 1}
    listing = (await client.get(f"/api/elements?project_id={project_id}&keyword=已更新", headers=headers)).json()
    assert listing["total"] == 1


async def test_element_excel_import_is_atomic_and_checks_creator(client: AsyncClient):
    """任一行失败整批回滚，且不允许更新其他用户创建的元素。"""
    headers, project_id = await _setup(client)
    before = await client.post(
        "/api/elements",
        headers=headers,
        json={"project_id": project_id, "name": "不应写入", "locator_type": "id", "locator_value": "before"},
    )
    assert before.status_code == 201
    response = await client.post(
        f"/api/projects/{project_id}/elements/import",
        headers=headers,
        files={"file": ("bad.xlsx", _element_xlsx([
            [None, project_id, "元素项目", "有效行", "首页", "android", "all", "id", "ok", None, None],
            [None, project_id, "元素项目", "错误行", "首页", "android", "all", "smart", "不得填写", "{bad", None],
        ]), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert response.status_code == 422
    detail = response.json()["detail"]
    assert detail["code"] == "ELEMENT_IMPORT_INVALID"
    assert any(error["row"] == 3 for error in detail["errors"])
    listing = (await client.get(f"/api/elements?project_id={project_id}&keyword=有效行", headers=headers)).json()
    assert listing["total"] == 0


def test_element_excel_large_smart_config_round_trip():
    """超长 smart 配置通过独立分片工作表导出后仍可完整导入。"""
    condition = {"attribute": "text", "operator": "equals", "value": "x" * 200}
    config = {
        "version": 1,
        "alternatives": [{"target": [condition.copy() for _ in range(20)]} for _ in range(10)],
        "search": {"scroll": True, "direction": "up", "max_swipes": 8, "duration_ms": 500, "settle_ms": 300},
        "selection": {"policy": "unique", "index": None},
    }
    element = SimpleNamespace(
        id=1,
        name="超长配置",
        page_name=None,
        platform="android",
        scope="all",
        locator_type="smart",
        locator_value=None,
        locator_config=config,
        description=None,
    )
    project = SimpleNamespace(id=1, name="项目")

    exported = build_export([(element, project)])
    workbook = load_workbook(BytesIO(exported), read_only=True, data_only=False)
    assert CONFIG_SHEET_TITLE in workbook.sheetnames
    main_value = workbook["元素"]["J2"].value
    assert isinstance(main_value, str) and main_value.startswith(CONFIG_REF_PREFIX)
    assert len(main_value) <= EXCEL_CELL_MAX_CHARS
    assert all(
        len(chunk) <= CONFIG_CHUNK_SIZE <= EXCEL_CELL_MAX_CHARS
        for row in workbook[CONFIG_SHEET_TITLE].iter_rows(min_row=2, values_only=True)
        if isinstance(chunk := row[2], str)
    )
    workbook.close()

    rows, errors = parse_import(exported, project_id=1)
    assert errors == []
    assert len(rows) == 1
    assert rows[0].data["locator_config"] == config


def test_element_excel_import_rejects_nonblank_rows_after_limit():
    """超过 2,000 条的数据不能在限制行之后被静默忽略。"""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "元素"
    sheet.append(list(ELEMENT_HEADERS))
    sheet.cell(row=2, column=4, value="第一条")
    sheet.cell(row=2, column=8, value="id")
    sheet.cell(row=2, column=9, value="first")
    sheet.cell(row=2003, column=4, value="超限条目")
    sheet.cell(row=2003, column=8, value="id")
    sheet.cell(row=2003, column=9, value="overflow")
    output = BytesIO()
    workbook.save(output)

    rows, errors = parse_import(output.getvalue(), project_id=1)
    assert len(rows) == 1
    assert any(error["row"] == 2003 for error in errors)


def test_element_excel_config_chunks_preserve_json_whitespace():
    """配置分片首尾空白属于 JSON 字符串内容时不能被导入清洗。"""
    config_text = '{"value":"a "}'
    reference = CONFIG_REF_PREFIX + sha256(config_text.encode("utf-8")).hexdigest()
    workbook = Workbook()
    data = workbook.active
    data.title = "元素"
    data.append(list(ELEMENT_HEADERS))
    data.append([None, 1, "项目", "带空格", None, "android", "all", "smart", None, reference, None])
    config_sheet = workbook.create_sheet(CONFIG_SHEET_TITLE)
    config_sheet.append(list(CONFIG_HEADERS))
    config_sheet.append([reference, 1, '{"value":"a '])
    config_sheet.append([reference, 2, '"}'])
    output = BytesIO()
    workbook.save(output)

    rows, errors = parse_import(output.getvalue(), project_id=1)
    assert errors == []
    assert rows[0].data["locator_config"] == {"value": "a "}


async def test_element_excel_import_validates_database_text_lengths(client: AsyncClient):
    """Excel 导入应在行级校验阶段拒绝超出数据库字段长度的文本。"""
    headers, project_id = await _setup(client)
    response = await client.post(
        f"/api/projects/{project_id}/elements/import",
        headers=headers,
        files={
            "file": (
                "length.xlsx",
                _element_xlsx([
                    [None, project_id, "元素项目", "页面过长", "x" * 256, "android", "all", "id", "page", None, None],
                    [None, project_id, "元素项目", "范围过长", "首页", "android", "s" * 101, "id", "scope", None, None],
                ]),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert response.status_code == 422
    detail = response.json()["detail"]
    assert detail["code"] == "ELEMENT_IMPORT_INVALID"
    assert {error["row"] for error in detail["errors"]} == {2, 3}
